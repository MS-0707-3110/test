"""
Excel ファイル管理モジュール
山梨県土地情報を .xlsx 形式で保存・更新する
"""

import os
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter

EXCEL_PATH = os.path.join(os.path.dirname(__file__), 'data', '山梨土地情報.xlsx')

# 列定義: (ヘッダー名, 列幅)
COLUMNS = [
    ('No.',        6),
    ('住所',       40),
    ('価格',       18),
    ('面積(m²)',   14),
    ('面積(坪)',   12),
    ('情報源',     12),
    ('参考URL',    60),
    ('取得日',     18),
]

# スタイル定数
HEADER_FILL  = PatternFill("solid", fgColor="1F6B4E")   # 濃い緑
ALT_ROW_FILL = PatternFill("solid", fgColor="E8F5E9")   # 薄い緑
HEADER_FONT  = Font(name="Meiryo", bold=True, color="FFFFFF", size=11)
DATA_FONT    = Font(name="Meiryo", size=10)
URL_FONT     = Font(name="Meiryo", size=10, color="1155CC", underline="single")
THIN_BORDER  = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'),  bottom=Side(style='thin'),
)


def update_excel(listings: list[dict]) -> str:
    """
    listings: scraper が返す dict のリスト
    重複排除・ソート後に Excel を上書き保存する。
    戻り値: 保存したファイルパス
    """
    os.makedirs(os.path.dirname(EXCEL_PATH), exist_ok=True)

    # 既存ファイルがあれば読み込み、なければ新規作成
    if os.path.exists(EXCEL_PATH):
        wb = load_workbook(EXCEL_PATH)
        ws = wb.active
        existing = _load_existing_rows(ws)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "土地情報"
        existing = {}

    # 新規データをマージ（URLをキーに重複排除）
    now_str = datetime.now().strftime('%Y-%m-%d %H:%M')
    for item in listings:
        url = item.get('参考URL', '')
        if not url:
            continue
        if url not in existing:
            existing[url] = {
                '住所':    item.get('住所', ''),
                '価格':    _format_price(item),
                '面積sqm': item.get('面積_sqm', ''),
                '情報源':  item.get('情報源', ''),
                '参考URL': url,
                '取得日':  now_str,
            }

    # シートを全消去して書き直す
    ws.delete_rows(1, ws.max_row + 1)
    _write_header(ws)
    _write_data(ws, list(existing.values()))

    wb.save(EXCEL_PATH)
    return EXCEL_PATH


def _load_existing_rows(ws) -> dict:
    """既存シートからURL→行データの辞書を作成"""
    existing = {}
    headers = [cell.value for cell in ws[1]]
    try:
        url_col = headers.index('参考URL')
    except ValueError:
        return existing

    for row in ws.iter_rows(min_row=2, values_only=True):
        url = row[url_col] if url_col < len(row) else None
        if url:
            existing[url] = {
                '住所':    row[headers.index('住所')] if '住所' in headers else '',
                '価格':    row[headers.index('価格')] if '価格' in headers else '',
                '面積sqm': _sqm_from_str(row[headers.index('面積(m²)')] if '面積(m²)' in headers else ''),
                '情報源':  row[headers.index('情報源')] if '情報源' in headers else '',
                '参考URL': url,
                '取得日':  row[headers.index('取得日')] if '取得日' in headers else '',
            }
    return existing


def _write_header(ws):
    for col_idx, (header, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font      = HEADER_FONT
        cell.fill      = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border    = THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = 'A2'


def _write_data(ws, rows: list[dict]):
    # 価格でソート（安い順）
    rows.sort(key=lambda x: _price_sort_key(x.get('価格', '')))

    for row_idx, item in enumerate(rows, start=2):
        fill = ALT_ROW_FILL if row_idx % 2 == 0 else None
        sqm  = item.get('面積sqm') or ''
        tsubo = f"{float(sqm) / 3.30579:.1f}" if sqm else ''

        values = [
            row_idx - 1,
            item.get('住所', ''),
            item.get('価格', ''),
            sqm,
            tsubo,
            item.get('情報源', ''),
            item.get('参考URL', ''),
            item.get('取得日', ''),
        ]

        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border    = THIN_BORDER
            cell.alignment = Alignment(vertical='center', wrap_text=(col_idx == 2))

            # URL列はハイパーリンクスタイル
            if col_idx == 7 and value:
                cell.font = URL_FONT
                cell.hyperlink = value
            else:
                cell.font = DATA_FONT

            if fill:
                cell.fill = fill

        ws.row_dimensions[row_idx].height = 18

    # 最終更新日時を A1 コメント的な位置（最終行の下）に記録
    last_row = ws.max_row + 2
    cell = ws.cell(row=last_row, column=1,
                   value=f"最終更新: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    cell.font = Font(name="Meiryo", size=9, color="888888", italic=True)


def _format_price(item: dict) -> str:
    man = item.get('価格_万円')
    if man:
        return f"{int(man):,}万円"
    return item.get('価格テキスト', '')


def _price_sort_key(price_str: str) -> float:
    """ソート用に価格文字列を数値化"""
    m = __import__('re').search(r'([\d,]+)', str(price_str).replace(',', ''))
    return float(m.group(1)) if m else 9_999_999


def _sqm_from_str(val) -> float | None:
    if val is None:
        return None
    try:
        return float(str(val).replace('m²', '').strip())
    except ValueError:
        return None
